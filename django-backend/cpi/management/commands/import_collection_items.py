from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from cpi.models import DimOutlet, CpiCollectionItem, CpiBasket
from basket_matcher import BasketMatcher


EXCEL_FILE = r"F:\Maldives-CPI-System\docs\Other_Docs\Male July final 2025\Male price list.xlsx"


def get_latest_price(ws, row):
    for col in range(ws.max_column, 9, -1):
        price = ws.cell(row=row, column=col).value
        month = ws.cell(row=1, column=col).value

        if price not in (None, ""):
            return price, month

    return None, None


def get_start_row():
    last_item = CpiCollectionItem.objects.order_by("-sort_order").first()

    if last_item and last_item.sort_order:
        return int(last_item.sort_order) + 2

    return 2


class Command(BaseCommand):

    help = "Import CPI collection items"

    def handle(self, *args, **options):

        self.stdout.write("Opening Excel file...")

        wb = load_workbook(EXCEL_FILE, data_only=True)
        ws = wb["by outlet"]

        matcher = BasketMatcher()

        self.stdout.write(f"Worksheet : {ws.title}")
        self.stdout.write(f"Rows      : {ws.max_row}")
        self.stdout.write(f"Columns   : {ws.max_column}")
        self.stdout.write("-" * 50)

        start_row = get_start_row()
        self.stdout.write(f"Starting from Excel row: {start_row}")

        for row in range(start_row, ws.max_row + 1):

            item_code = ws.cell(row=row, column=1).value
            product_name = ws.cell(row=row, column=2).value
            outlet_name = ws.cell(row=row, column=5).value
            brand = ws.cell(row=row, column=6).value
            unit = ws.cell(row=row, column=7).value
            excel_product_code = ws.cell(row=row, column=8).value
            specification = ws.cell(row=row, column=9).value

            if not product_name:
                continue

            outlet_name_clean = outlet_name.strip() if outlet_name else ""

            outlet = DimOutlet.objects.filter(
                outlet_name__iexact=outlet_name_clean,
                is_active=True
            ).first()

            if not outlet:
                raise Exception(f"Row {row}: Outlet not found: {outlet_name}")

            basket = CpiBasket.objects.filter(
                basket_code=str(item_code).zfill(3),
                is_basket_item=True,
                is_active=True
            ).first()

            if not basket:
                raise Exception(
                    f"Row {row}: Basket code not found: {item_code} - {product_name}"
                )
            

            name_score = matcher.similarity(product_name, basket.basket_item_name)

            if name_score < 0.75:
                raise Exception(
                    f"\nBasket name mismatch!\n\n"
                    f"Row            : {row}\n"
                    f"Outlet         : {outlet.outlet_code} - {outlet.outlet_name}\n"
                    f"Excel code     : {str(item_code).zfill(3)}\n"
                    f"Basket code    : {basket.basket_code}\n"
                    f"Excel product  : {product_name}\n"
                    f"Basket product : {basket.basket_item_name}\n"
                    f"Name score     : {round(name_score * 100)}%\n"
                )

            last_price, last_price_month = get_latest_price(ws, row)

            collection_item, created = CpiCollectionItem.objects.update_or_create(
                outlet=outlet,
                legacy_item_code=item_code,
                specification_no=specification,
                defaults={
                    "basket_item": basket,
                    "product_name": product_name,
                    "matched_basket_name": basket.basket_item_name,
                    "brand": brand,
                    "unit": unit,
                    "sort_order": row - 1,
                    "excel_product_code": excel_product_code,
                    "excel_specification": specification,
                    "last_price": last_price,
                    "last_price_month": last_price_month.strftime("%Y-%m") if last_price_month else None,
                }
            )

            self.stdout.write(
                f"{'Created' if created else 'Updated'}: "
                f"{collection_item.product_code} | {product_name}"
            )

        self.stdout.write("-" * 50)
        self.stdout.write("Import finished.")