from datetime import datetime

from django.core.management.base import BaseCommand
from django.utils import timezone
from openpyxl import load_workbook

from cpi.models import (
    CpiMethodology,
    CpiRound,
    CpiVisit,
    CpiCollectionItem,
    FactPrice,
)


EXCEL_FILE = (
    r"F:\Maldives-CPI-System\docs\Other_Docs"
    r"\Male July final 2025\Male price list.xlsx"
)

METHODOLOGY_CODE = "CPI2022"


def get_excel_month(value):
    """
    Convert the Excel month header into year and month.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.year, value.month

    if hasattr(value, "year") and hasattr(value, "month"):
        return value.year, value.month

    value = str(value).strip()

    formats = [
        "%Y-%m",
        "%Y-%m-%d",
        "%b-%y",
        "%b %y",
        "%b-%Y",
        "%b %Y",
        "%B-%y",
        "%B %y",
        "%B-%Y",
        "%B %Y",
    ]

    for date_format in formats:
        try:
            parsed_date = datetime.strptime(value, date_format)
            return parsed_date.year, parsed_date.month
        except ValueError:
            continue

    return None


class Command(BaseCommand):

    help = "Import historical CPI prices from Male price list"

    def handle(self, *args, **options):

        self.stdout.write("Opening Excel file...")

        wb = load_workbook(
            EXCEL_FILE,
            data_only=True
        )

        ws = wb["by outlet"]

        methodology = CpiMethodology.objects.get(
            methodology_code=METHODOLOGY_CODE
        )

        self.stdout.write(f"Worksheet : {ws.title}")
        self.stdout.write(f"Rows      : {ws.max_row}")
        self.stdout.write(f"Columns   : {ws.max_column}")
        self.stdout.write("-" * 60)

        # ---------------------------------------------------
        # Read month columns
        # ---------------------------------------------------

        month_columns = []

        for col in range(10, ws.max_column + 1):

            month_value = ws.cell(
                row=1,
                column=col
            ).value

            month_info = get_excel_month(month_value)

            if not month_info:
                self.stdout.write(
                    self.style.WARNING(
                        f"Skipping column {col}: "
                        f"Invalid month header '{month_value}'"
                    )
                )
                continue

            year, month = month_info

            month_columns.append(
                {
                    "column": col,
                    "year": year,
                    "month": month,
                }
            )

        self.stdout.write(
            f"Valid month columns: {len(month_columns)}"
        )

        # ---------------------------------------------------
        # Create/find CPI rounds
        # ---------------------------------------------------

        rounds = {}

        for month_info in month_columns:

            year = month_info["year"]
            month = month_info["month"]

            round_obj, created = CpiRound.objects.get_or_create(
                survey_year=year,
                survey_month=month,
                defaults={
                    "methodology": methodology,
                    "round_status": "LOCKED",
                    "created_at": timezone.now(),
                }
            )

            rounds[(year, month)] = round_obj

            self.stdout.write(
                f"{'Created' if created else 'Found'} round: "
                f"{year}-{month:02d}"
            )

        # ---------------------------------------------------
        # Import historical prices
        # ---------------------------------------------------

        created_prices = 0
        updated_prices = 0
        skipped_blank = 0

        visit_cache = {}

        for row in range(2, ws.max_row + 1):

            item_code = ws.cell(row=row, column=1).value
            product_name = ws.cell(row=row, column=2).value
            outlet_name = ws.cell(row=row, column=5).value
            specification = ws.cell(row=row, column=9).value

            if not product_name:
                continue

            outlet_name_clean = (
                outlet_name.strip()
                if outlet_name
                else ""
            )

            collection_item = (
                CpiCollectionItem.objects
                .filter(
                    sort_order=row - 1
                )
                .first()
            )
            if not collection_item:
                raise Exception(
                    f"\nCollection item not found!\n\n"
                    f"Excel row      : {row}\n"
                    f"Sort order     : {row - 1}\n"
                    f"Outlet         : {outlet_name}\n"
                    f"Item code      : {item_code}\n"
                    f"Specification  : {specification}\n"
                    f"Product        : {product_name}\n"
                )

            excel_code = str(item_code).zfill(3)
            collection_code = str(collection_item.legacy_item_code).zfill(3)

            if excel_code != collection_code:
                raise Exception(
                    f"\nCollection item code mismatch!\n\n"
                    f"Excel row          : {row}\n"
                    f"Sort order         : {row - 1}\n"
                    f"Excel code         : {excel_code}\n"
                    f"Collection code    : {collection_code}\n"
                    f"Excel product      : {product_name}\n"
                    f"Collection product : {collection_item.product_name}\n"
                )

            excel_outlet = outlet_name_clean.lower()
            collection_outlet = collection_item.outlet.outlet_name.strip().lower()

            if excel_outlet != collection_outlet:
                raise Exception(
                    f"\nCollection item outlet mismatch!\n\n"
                    f"Excel row          : {row}\n"
                    f"Sort order         : {row - 1}\n"
                    f"Excel outlet       : {outlet_name}\n"
                    f"Collection outlet  : {collection_item.outlet.outlet_name}\n"
                    f"Excel product      : {product_name}\n"
                    f"Collection product : {collection_item.product_name}\n"
                )

            previous_price = None

            for month_info in month_columns:

                col = month_info["column"]
                year = month_info["year"]
                month = month_info["month"]

                price = ws.cell(
                    row=row,
                    column=col
                ).value

                if price in (None, ""):
                    skipped_blank += 1
                    continue

                round_obj = rounds[(year, month)]

                visit_key = (
                    round_obj.round_id,
                    collection_item.outlet_id,
                )

                if visit_key not in visit_cache:

                    visit, visit_created = (
                        CpiVisit.objects.get_or_create(
                            round=round_obj,
                            outlet=collection_item.outlet,
                            defaults={
                                "visit_status": "COMPLETED",
                                "sync_status": "SYNCED",
                                "created_on_device": False,
                                "last_modified_at": timezone.now(),
                                "is_deleted": False,
                                "created_at": timezone.now(),
                                "review_status": "APPROVED",
                                "reviewed_at": timezone.now(),
                            }
                        )
                    )

                    visit_cache[visit_key] = visit

                visit = visit_cache[visit_key]

                fact_price, created = (
                    FactPrice.objects.update_or_create(
                        visit=visit,
                        collection_item=collection_item,
                        defaults={
                            "availability": "AVAILABLE",
                            "observed_price": price,
                            "last_month_price": previous_price,
                            "remarks": "Imported historical CPI price",
                            "review_status": "APPROVED",
                            "reviewed_at": timezone.now(),
                            "is_outlier": False,
                            "quote_status": "ACCEPTED",
                            "sync_status": "SYNCED",
                            "created_on_device": False,
                            "server_received_at": timezone.now(),
                            "last_modified_at": timezone.now(),
                            "is_deleted": False,
                            "created_at": timezone.now(),
                            "updated_at": timezone.now(),
                        }
                    )
                )

                if created:
                    created_prices += 1
                else:
                    updated_prices += 1

                previous_price = price

            self.stdout.write(
                f"Processed row {row}: "
                f"{collection_item.product_code} | "
                f"{product_name}"
            )

        # ---------------------------------------------------
        # Summary
        # ---------------------------------------------------

        self.stdout.write("-" * 60)

        self.stdout.write(
            self.style.SUCCESS(
                "Historical price import finished."
            )
        )

        self.stdout.write(
            f"Created prices : {created_prices}"
        )

        self.stdout.write(
            f"Updated prices : {updated_prices}"
        )

        self.stdout.write(
            f"Blank prices   : {skipped_blank}"
        )