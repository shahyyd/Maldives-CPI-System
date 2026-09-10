import os
import sys
import django
from openpyxl import load_workbook

print("Current folder:", os.getcwd())
print("Python path:", sys.path[0])

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpi_admin.settings")
django.setup()

from cpi.models import CpiMethodology, CpiBasket


EXCEL_FILE = r"F:\Maldives-CPI-System\docs\CPI Basket.xlsx"
SHEET_NAME = "CPI basket "
METHODOLOGY_CODE = "CPI2022"


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value if value else None


def clean_level(value):
    if value is None or str(value).strip() == "":
        return None
    return int(value)


def main():
    methodology = CpiMethodology.objects.get(methodology_code=METHODOLOGY_CODE)

    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]

    #CpiBasket.objects.filter(methodology=methodology).delete()

    latest_by_level = {}
    created = 0
    skipped = 0
    display_order = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        basket_code = clean(row[0])
        level = clean_level(row[1])
        coicop_2016 = clean(row[2])
        description = clean(row[3])

        if not basket_code or not description:
            skipped += 1
            continue

        display_order += 1
        is_basket_item = level is None

        if is_basket_item:
            parent_basket = latest_by_level[max(latest_by_level.keys())] if latest_by_level else None
        else:
            parent_basket = latest_by_level.get(level - 1)

        # Current hierarchy context
        level_2 = latest_by_level.get(2)
        level_3 = latest_by_level.get(3)
        level_4 = latest_by_level.get(4)
        level_5 = latest_by_level.get(5)

        obj = CpiBasket.objects.create(
            methodology=methodology,
            parent_basket=parent_basket,
            basket_code=basket_code,
            level=level,
            coicop_2016=coicop_2016,
            description=description,
            is_basket_item=is_basket_item,
            display_order=display_order,

            coicop_division_code=level_2.basket_code if level_2 else None,
            coicop_division_name=level_2.description if level_2 else None,

            coicop_group_code=level_3.basket_code if level_3 else None,
            coicop_group_name=level_3.description if level_3 else None,

            coicop_class_code=level_4.basket_code if level_4 else None,
            coicop_class_name=level_4.description if level_4 else None,

            coicop_subclass_code=level_5.basket_code if level_5 else None,
            coicop_subclass_name=level_5.description if level_5 else None,

            basket_item_code=basket_code if is_basket_item else None,
            basket_item_name=description if is_basket_item else None,

            is_active=True,
        )

        if not is_basket_item:
            latest_by_level[level] = obj

            for old_level in list(latest_by_level.keys()):
                if old_level > level:
                    del latest_by_level[old_level]

        created += 1

    print("CPI Basket import completed.")
    print(f"Created: {created}")
    print(f"Skipped empty rows: {skipped}")


if __name__ == "__main__":
    main()