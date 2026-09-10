from collections import Counter
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from cpi.models import CpiCollectionItem


DEFAULT_EXCEL_FILE = (
    r"F:\Maldives-CPI-System\docs\All Outlets.xlsx"
)

HEADER_ROW = 5
DATA_START_ROW = HEADER_ROW + 1

PRODUCT_NAME_COLUMN = 1
SPECIFICATION_COLUMN = 2
EXCEL_PRODUCT_CODE_COLUMN = 3

HEADER_CODE_VALUES = {
    "product code",
    "product_code",
    "excel product code",
    "excel_product_code",
}


def clean_text(value):
    """
    Convert an Excel cell value into clean text.

    Returns None when the cell is empty or contains only spaces.
    """
    if value is None:
        return None

    text = str(value).strip()
    return text if text else None


def clean_excel_code(value):
    """
    Normalize an Excel product code.

    Examples:
        123       -> "123"
        123.0     -> "123"
        " 123 "   -> "123"
        None      -> None

    Codes containing leading zeros should preferably be stored as text
    in Excel, for example "00123".
    """
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return clean_text(value)


def is_header_code(value):
    """
    Return True when a cell contains a repeated column heading rather
    than an actual Excel product code.
    """
    if not value:
        return False

    normalized = " ".join(str(value).strip().lower().split())
    return normalized in HEADER_CODE_VALUES


class Command(BaseCommand):
    help = (
        "Update collection-item product specifications from "
        "the All Outlets workbook. Excel product codes that occur "
        "more than once are skipped for manual review."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            dest="excel_file",
            default=DEFAULT_EXCEL_FILE,
            help="Full path to the All Outlets Excel workbook.",
        )

        parser.add_argument(
            "--sheet",
            dest="sheet_name",
            help=(
                "Process only one worksheet. "
                "If omitted, all worksheets are processed."
            ),
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
            dest="dry_run",
            help="Validate and report changes without updating the database.",
        )

    def handle(self, *args, **options):
        excel_file = Path(options["excel_file"])
        sheet_name = options.get("sheet_name")
        dry_run = options["dry_run"]

        if not excel_file.exists():
            raise CommandError(
                f"Excel workbook not found:\n{excel_file}"
            )

        if dry_run:
            self.stdout.write(
                self.style.WARNING(
                    "DRY RUN MODE: no database records will be updated."
                )
            )

        self.stdout.write(f"Opening workbook: {excel_file}")

        try:
            workbook = load_workbook(
                filename=excel_file,
                data_only=True,
                read_only=True,
            )
        except Exception as exc:
            raise CommandError(
                f"Unable to open workbook: {exc}"
            ) from exc

        try:
            worksheets = self.get_worksheets(
                workbook=workbook,
                sheet_name=sheet_name,
            )

            self.stdout.write(
                "Counting Excel product-code occurrences..."
            )

            code_counts = self.count_workbook_codes(worksheets)
            duplicate_codes = {
                code
                for code, count in code_counts.items()
                if count > 1
            }

            if duplicate_codes:
                self.stdout.write(
                    self.style.WARNING(
                        f"{len(duplicate_codes):,} duplicate Excel "
                        "product code(s) will be skipped."
                    )
                )

            collection_items = self.build_collection_item_lookup()

            statistics = {
                "worksheets": 0,
                "rows_read": 0,
                "rows_without_code": 0,
                "rows_without_specification": 0,
                "header_rows_skipped": 0,
                "duplicate_rows_skipped": 0,
                "duplicate_codes_skipped": len(duplicate_codes),
                "codes_not_found": 0,
                "unchanged": 0,
                "to_update": 0,
            }

            pending_updates = []
            not_found_records = []
            duplicate_records = []

            for worksheet in worksheets:
                statistics["worksheets"] += 1

                self.stdout.write(
                    f"Processing sheet: {worksheet.title}"
                )

                for row_number, row_values in enumerate(
                    worksheet.iter_rows(
                        min_row=DATA_START_ROW,
                        values_only=True,
                    ),
                    start=DATA_START_ROW,
                ):
                    statistics["rows_read"] += 1

                    product_name = clean_text(
                        self.get_column_value(
                            row_values,
                            PRODUCT_NAME_COLUMN,
                        )
                    )

                    specification_text = clean_text(
                        self.get_column_value(
                            row_values,
                            SPECIFICATION_COLUMN,
                        )
                    )

                    excel_product_code = clean_excel_code(
                        self.get_column_value(
                            row_values,
                            EXCEL_PRODUCT_CODE_COLUMN,
                        )
                    )

                    if not excel_product_code:
                        statistics["rows_without_code"] += 1
                        continue

                    if is_header_code(excel_product_code):
                        statistics["header_rows_skipped"] += 1
                        continue

                    if not specification_text:
                        statistics[
                            "rows_without_specification"
                        ] += 1
                        continue

                    if excel_product_code in duplicate_codes:
                        statistics["duplicate_rows_skipped"] += 1

                        duplicate_records.append(
                            {
                                "sheet": worksheet.title,
                                "row": row_number,
                                "code": excel_product_code,
                                "product_name": product_name,
                                "specification": specification_text,
                            }
                        )
                        continue

                    collection_item = collection_items.get(
                        excel_product_code
                    )

                    if collection_item is None:
                        statistics["codes_not_found"] += 1

                        not_found_records.append(
                            {
                                "sheet": worksheet.title,
                                "row": row_number,
                                "code": excel_product_code,
                                "product_name": product_name,
                            }
                        )
                        continue

                    specification_changed = (
                        clean_text(
                            collection_item.product_specification
                        )
                        != specification_text
                    )

                    excel_specification_changed = (
                        clean_text(
                            collection_item.excel_specification
                        )
                        != specification_text
                    )

                    if not (
                        specification_changed
                        or excel_specification_changed
                    ):
                        statistics["unchanged"] += 1
                        continue

                    collection_item.product_specification = (
                        specification_text
                    )
                    collection_item.excel_specification = (
                        specification_text
                    )

                    pending_updates.append(collection_item)
                    statistics["to_update"] += 1

            self.print_duplicate_records(duplicate_records)
            self.print_not_found_records(not_found_records)

            if pending_updates and not dry_run:
                self.save_updates(pending_updates)

            self.print_summary(
                statistics=statistics,
                dry_run=dry_run,
            )

        finally:
            workbook.close()

    def get_worksheets(self, workbook, sheet_name):
        """
        Return either the requested worksheet or all worksheets.
        """
        if not sheet_name:
            return workbook.worksheets

        if sheet_name not in workbook.sheetnames:
            available_sheets = ", ".join(workbook.sheetnames)

            raise CommandError(
                f'Worksheet "{sheet_name}" was not found.\n'
                f"Available worksheets: {available_sheets}"
            )

        return [workbook[sheet_name]]

    def count_workbook_codes(self, worksheets):
        """
        Count valid Excel product-code occurrences across the selected
        worksheets. Repeated header rows and blank cells are excluded.
        """
        code_counts = Counter()

        for worksheet in worksheets:
            for row_values in worksheet.iter_rows(
                min_row=DATA_START_ROW,
                values_only=True,
            ):
                excel_product_code = clean_excel_code(
                    self.get_column_value(
                        row_values,
                        EXCEL_PRODUCT_CODE_COLUMN,
                    )
                )

                if not excel_product_code:
                    continue

                if is_header_code(excel_product_code):
                    continue

                code_counts[excel_product_code] += 1

        return code_counts

    def build_collection_item_lookup(self):
        """
        Build a dictionary keyed by excel_product_code.

        The current database is expected to contain no duplicate
        nonblank Excel product codes. This validation protects the
        update from modifying an ambiguous database record.
        """
        self.stdout.write(
            "Loading collection-item codes from the database..."
        )

        lookup = {}
        duplicate_codes = []

        queryset = (
            CpiCollectionItem.objects
            .exclude(excel_product_code__isnull=True)
            .exclude(excel_product_code="")
            .only(
                "collection_item_id",
                "excel_product_code",
                "product_name",
                "product_specification",
                "excel_specification",
            )
            .iterator(chunk_size=2000)
        )

        for collection_item in queryset:
            code = clean_excel_code(
                collection_item.excel_product_code
            )

            if not code:
                continue

            if code in lookup:
                duplicate_codes.append(code)
                continue

            lookup[code] = collection_item

        if duplicate_codes:
            duplicate_list = ", ".join(
                sorted(set(duplicate_codes))[:20]
            )

            raise CommandError(
                "Duplicate excel_product_code values were found "
                f"in the database:\n{duplicate_list}"
            )

        self.stdout.write(
            f"Loaded {len(lookup):,} collection items."
        )

        return lookup

    def save_updates(self, pending_updates):
        """
        Save all changes as one protected database operation.
        """
        self.stdout.write(
            f"Updating {len(pending_updates):,} records..."
        )

        with transaction.atomic():
            CpiCollectionItem.objects.bulk_update(
                pending_updates,
                fields=[
                    "product_specification",
                    "excel_specification",
                ],
                batch_size=1000,
            )

    def print_duplicate_records(self, records):
        """
        Show duplicate workbook codes skipped for manual review.
        Detailed output is limited to avoid flooding the terminal.
        """
        if not records:
            return

        self.stdout.write("")
        self.stdout.write(
            self.style.WARNING(
                "Duplicate Excel product-code rows skipped:"
            )
        )

        display_limit = 50

        for record in records[:display_limit]:
            specification_preview = (
                record["specification"]
                .replace("\r", " ")
                .replace("\n", " | ")
            )

            if len(specification_preview) > 160:
                specification_preview = (
                    specification_preview[:157] + "..."
                )

            self.stdout.write(
                self.style.WARNING(
                    f"Sheet: {record['sheet']} | "
                    f"Row: {record['row']} | "
                    f"Code: {record['code']} | "
                    f"Product: {record['product_name'] or '-'} | "
                    f"Specification: {specification_preview}"
                )
            )

        remaining = len(records) - display_limit

        if remaining > 0:
            self.stdout.write(
                self.style.WARNING(
                    f"...and {remaining:,} additional duplicate rows."
                )
            )

    def print_not_found_records(self, records):
        """
        Display unmatched codes, limiting detailed output to avoid
        flooding the terminal.
        """
        if not records:
            return

        self.stdout.write("")
        self.stdout.write(
            self.style.WARNING(
                "Excel product codes not found in the database:"
            )
        )

        display_limit = 50

        for record in records[:display_limit]:
            self.stdout.write(
                self.style.WARNING(
                    f"Sheet: {record['sheet']} | "
                    f"Row: {record['row']} | "
                    f"Code: {record['code']} | "
                    f"Product: {record['product_name'] or '-'}"
                )
            )

        remaining = len(records) - display_limit

        if remaining > 0:
            self.stdout.write(
                self.style.WARNING(
                    f"...and {remaining:,} additional unmatched codes."
                )
            )

    def print_summary(self, statistics, dry_run):
        self.stdout.write("")
        self.stdout.write("-" * 60)

        if dry_run:
            title = "Product specification dry run finished."
        else:
            title = "Product specification update finished."

        self.stdout.write(self.style.SUCCESS(title))

        action_label = (
            "Records to update"
            if dry_run
            else "Records updated"
        )

        self.stdout.write(
            f"Worksheets processed        : "
            f"{statistics['worksheets']:,}"
        )
        self.stdout.write(
            f"Rows read                   : "
            f"{statistics['rows_read']:,}"
        )
        self.stdout.write(
            f"{action_label:<28}: "
            f"{statistics['to_update']:,}"
        )
        self.stdout.write(
            f"Already unchanged           : "
            f"{statistics['unchanged']:,}"
        )
        self.stdout.write(
            f"Rows without product code   : "
            f"{statistics['rows_without_code']:,}"
        )
        self.stdout.write(
            f"Rows without specification  : "
            f"{statistics['rows_without_specification']:,}"
        )
        self.stdout.write(
            f"Header rows skipped         : "
            f"{statistics['header_rows_skipped']:,}"
        )
        self.stdout.write(
            f"Duplicate codes skipped     : "
            f"{statistics['duplicate_codes_skipped']:,}"
        )
        self.stdout.write(
            f"Duplicate rows skipped      : "
            f"{statistics['duplicate_rows_skipped']:,}"
        )
        self.stdout.write(
            f"Codes not found             : "
            f"{statistics['codes_not_found']:,}"
        )

    @staticmethod
    def get_column_value(row_values, column_number):
        """
        Safely return a value from a row using a one-based Excel
        column number.
        """
        index = column_number - 1

        if index >= len(row_values):
            return None

        return row_values[index]