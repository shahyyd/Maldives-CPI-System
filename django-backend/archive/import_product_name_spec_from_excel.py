import os
import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpi_admin.settings")
django.setup()

from openpyxl import load_workbook
from django.db import transaction
from django.utils import timezone

from cpi.models import DimOutlet, DimItem, CpiCollectionItem


EXCEL_PATH = r"F:\CPI_CapiApplication\Other_Docs\Male July final 2025\First 2 weeks\Food forms\Food Form - with new products.xlsx"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def value_after_colon(value):
    text = clean(value)
    if ":" in text:
        return text.split(":", 1)[1].strip()
    return text


def next_specification_no(outlet, item):
    last_count = CpiCollectionItem.objects.filter(
        outlet=outlet,
        item=item
    ).count()
    return f"{last_count + 1:03d}"


wb = load_workbook(EXCEL_PATH, data_only=True)

updated_items = 0
created_items = 0
created_dim_items = 0
missing_outlets = []
skipped_rows = 0

with transaction.atomic():

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        outlet_name = value_after_colon(ws["A2"].value)

        if not outlet_name:
            continue

        if outlet_name.upper() in ["PRODUCT NAME", "ITEMS", "OTHERS", "BOOKSHOP"]:
            continue

        outlet = DimOutlet.objects.filter(excel_sheet_name=sheet_name).first()

        if not outlet:
            outlet = DimOutlet.objects.filter(outlet_name__iexact=outlet_name).first()

        if not outlet:
            missing_outlets.append(f"{sheet_name} / {outlet_name}")
            continue

        print(f"Processing: {outlet.outlet_code} - {outlet.outlet_name}")

        sort_order = 0

        for row in range(6, ws.max_row + 1):

            product_name = clean(ws.cell(row=row, column=1).value)
            product_specification = clean(ws.cell(row=row, column=2).value)

            if not product_name or not product_specification:
                skipped_rows += 1
                continue

            if product_name.upper() in ["PRODUCT NAME", "ITEMS", "OTHERS"]:
                skipped_rows += 1
                continue

            sort_order += 1

            collection_item = CpiCollectionItem.objects.filter(
                outlet=outlet,
                sort_order=sort_order
            ).first()

            dim_item, dim_created = DimItem.objects.get_or_create(
                item_name=product_name,
                defaults={
                    "is_active": True,
                }
            )

            if dim_created:
                created_dim_items += 1
                dim_item.save()

            if collection_item:
                collection_item.product_name = product_name
                collection_item.product_specification = product_specification
                collection_item.sort_order = sort_order
                collection_item.save()
                updated_items += 1

            else:
                collection_item = CpiCollectionItem.objects.create(
                    outlet=outlet,
                    item=dim_item,
                    specification_no=next_specification_no(outlet, dim_item),
                    product_name=product_name,
                    product_specification=product_specification,
                    sort_order=sort_order,
                    is_active=True,
                    valid_from=timezone.now().date(),
                    item_status="ACTIVE",
                )
                created_items += 1


print("\nPRODUCT IMPORT COMPLETED")
print(f"Updated CpiCollectionItem: {updated_items}")
print(f"Created CpiCollectionItem: {created_items}")
print(f"Created DimItem: {created_dim_items}")
print(f"Skipped rows: {skipped_rows}")

if missing_outlets:
    print("\nMissing outlets:")
    for x in sorted(set(missing_outlets)):
        print("-", x)