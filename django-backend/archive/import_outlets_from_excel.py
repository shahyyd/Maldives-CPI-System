import os
import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpi_admin.settings")
django.setup()

from openpyxl import load_workbook
from django.db import transaction
from django.utils import timezone

from cpi.models import DimOutlet, DimOutletType, DimIsland


EXCEL_PATH = r"F:\CPI_CapiApplication\Other_Docs\Male July final 2025\First 2 weeks\Food forms\Food Form - with new products.xlsx"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def value_after_colon(value):
    text = clean(value)
    if ":" in text:
        return text.split(":", 1)[1].strip()
    return text


def get_default_island():
    island = DimIsland.objects.filter(island_name__iexact="Maale").first()
    if not island:
        island = DimIsland.objects.filter(island_name__iexact="Male").first()
    if not island:
        island = DimIsland.objects.filter(island_name__icontains="Male").first()
    if not island:
        island = DimIsland.objects.first()
    return island


def find_island_from_address(address_text):
    text = clean(address_text).upper()

    if "HULHUMALE" in text or "HULHUMALÉ" in text:
        island = DimIsland.objects.filter(island_name__icontains="Hulhumale").first()
        if island:
            return island

    if "MAALE" in text or "MALE" in text or "K.MALE" in text or "K.MALE'" in text:
        island = DimIsland.objects.filter(island_name__iexact="Maale").first()
        if island:
            return island

    return get_default_island()


wb = load_workbook(EXCEL_PATH, data_only=True)

created_types = 0
created_outlets = 0
updated_outlets = 0
skipped_sheets = []

with transaction.atomic():

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        address_text = value_after_colon(ws["A1"].value)
        contact_no = value_after_colon(ws["B1"].value)

        outlet_name = value_after_colon(ws["A2"].value)
        contact_person = value_after_colon(ws["B2"].value)

        outlet_type_name = value_after_colon(ws["A3"].value)
        outlet_code_from_excel = value_after_colon(ws["A4"].value)

        if not outlet_name:
            skipped_sheets.append(sheet_name)
            continue

        if outlet_code_from_excel.upper() == "PRODUCT NAME" or outlet_name.upper() in [
            "PRODUCT NAME",
            "ITEMS",
            "OTHERS",
            "BOOKSHOP",
        ]:
            skipped_sheets.append(sheet_name)
            continue

        island = find_island_from_address(address_text)

        outlet_type = None
        if outlet_type_name:
            outlet_type, type_created = DimOutletType.objects.get_or_create(
                outlet_type_name=outlet_type_name,
                defaults={"broad_type": outlet_type_name},
            )

            if type_created:
                created_types += 1

        outlet = DimOutlet.objects.filter(outlet_name__iexact=outlet_name).first()

        if outlet:
            # IMPORTANT:
            # Do not overwrite existing outlet_code from Excel.
            # We already cleaned outlet codes manually to 4 digits.
            if not outlet.outlet_code or str(outlet.outlet_code).strip() == "":
                if outlet_code_from_excel and outlet_code_from_excel.isdigit():
                    outlet.outlet_code = outlet_code_from_excel.zfill(4)

            outlet.outlet_type = outlet_type
            outlet.address_text = address_text
            outlet.contact_no = contact_no
            outlet.contact_person = contact_person
            outlet.excel_sheet_name = sheet_name
            outlet.island = island
            outlet.updated_at = timezone.now()
            outlet.save()

            updated_outlets += 1
            print(f"Updated outlet: {outlet.outlet_name} -> {island.island_name}")

        else:
            # New outlets get NULL outlet_code first.
            # We assign unique 4-digit codes using SQL after import.
            outlet = DimOutlet.objects.create(
                outlet_code=None,
                outlet_name=outlet_name,
                outlet_type=outlet_type,
                address_text=address_text,
                contact_no=contact_no,
                contact_person=contact_person,
                excel_sheet_name=sheet_name,
                island=island,
                is_active=True,
                created_at=timezone.now(),
                updated_at=timezone.now(),
            )

            created_outlets += 1
            print(f"Created outlet: {outlet.outlet_name} -> {island.island_name}")


print("\nOUTLET IMPORT COMPLETED")
print(f"Outlet types created: {created_types}")
print(f"Outlets created: {created_outlets}")
print(f"Outlets updated: {updated_outlets}")

if skipped_sheets:
    print("\nSkipped sheets:")
    for sheet in skipped_sheets:
        print("-", sheet)