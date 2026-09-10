import os
import django
from decimal import Decimal, InvalidOperation

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpi_admin.settings")
django.setup()

from openpyxl import load_workbook
from django.db import transaction
from cpi.models import CpiCollectionItem


EXCEL_PATH = r"F:\CPI_CapiApplication\Other_Docs\Male July final 2025\Male price list.xlsx"
SHEET_NAME = "by outlet"


def clean(value):
    return str(value).strip() if value is not None else ""


def to_decimal(value):
    if value is None or value == "":
        return None

    if isinstance(value, (int, float)):
        price = Decimal(str(value))
        return price if price > 0 else None

    text = clean(value).replace(",", "")

    try:
        price = Decimal(text)
        return price if price > 0 else None
    except InvalidOperation:
        return None


def last_available_price_from_row(row):
    last_price = None

    # J onwards = price history
    for cell in row[9:]:
        price = to_decimal(cell.value)
        if price is not None:
            last_price = price

    return last_price


wb = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
ws = wb[SHEET_NAME]

updated = 0
skipped = 0
no_price = 0
not_found = 0

with transaction.atomic():

    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):

        product_name = clean(row[1].value) if len(row) > 1 else ""   # Column B
        outlet_name = clean(row[4].value) if len(row) > 4 else ""    # Column E

        if not product_name or not outlet_name:
            skipped += 1
            continue

        last_price = last_available_price_from_row(row)

        if last_price is None:
            no_price += 1
            continue

        qs = CpiCollectionItem.objects.filter(
            outlet__outlet_name__icontains=outlet_name,
            product_name__icontains=product_name,
        )

        if not qs.exists():
            qs = CpiCollectionItem.objects.filter(
                outlet__outlet_name__icontains=outlet_name,
                product_name__iexact=product_name,
            )

        if not qs.exists():
            qs = CpiCollectionItem.objects.filter(
                outlet__outlet_name__iexact=outlet_name,
                product_name__icontains=product_name,
            )

        count = qs.update(last_price=last_price)

        if count == 0:
            not_found += 1
        else:
            updated += count

        if row_number % 500 == 0:
            print(f"Processed rows: {row_number}, updated: {updated}")


print("\nRELAXED LAST PRICE UPDATE COMPLETED")
print(f"Updated collection item rows: {updated}")
print(f"Skipped rows: {skipped}")
print(f"Rows with no price: {no_price}")
print(f"Rows not found: {not_found}")