from openpyxl import load_workbook

print("******** NEW VERSION ********")

EXCEL_PATH = r"F:\CPI_CapiApplication\Other_Docs\Male July final 2025\First 2 weeks\Food forms\Food Form - with new products.xlsx"

wb = load_workbook(EXCEL_PATH, data_only=True)

sheet_name = "Nide Mart"
ws = wb[sheet_name]

print("SHEET:", sheet_name)
print("MAX ROWS:", ws.max_row)
print("MAX COLS:", ws.max_column)

print("\nROW 5 HEADERS:")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=5, column=col).value
    if value not in (None, ""):
        print(f"COL {col}: {value}")

print("\nROW 8 VALUES:")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=8, column=col).value
    if value not in (None, ""):
        print(f"COL {col}: {value}")

print("\nROW 9 VALUES:")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=9, column=col).value
    if value not in (None, ""):
        print(f"COL {col}: {value}")