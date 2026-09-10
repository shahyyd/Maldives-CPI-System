from openpyxl import load_workbook

EXCEL_PATH = r"F:\CPI_CapiApplication\Other_Docs\Male July final 2025\First 2 weeks\Food forms\Food Form - with new products.xlsx"

wb = load_workbook(EXCEL_PATH, data_only=True)

sheet_name = "Nide Mart"
ws = wb[sheet_name]

print("SHEET:", sheet_name)

for row in range(1, 30):
    a = ws.cell(row=row, column=1).value
    b = ws.cell(row=row, column=2).value
    c = ws.cell(row=row, column=3).value

    print(f"Row {row}")
    print("A =", repr(a))
    print("B =", repr(b))
    print("C =", repr(c))
    print("-" * 80)